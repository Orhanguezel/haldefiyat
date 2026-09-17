import json,csv,collections,pathlib
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
P=pathlib.Path(__file__).parent
x=json.loads((P/'live-evidence.json').read_text());cur={r['keys'][0]:r for r in x['gsc']['current']['query']};prev={r['keys'][0]:r for r in x['gsc']['previous']['query']};pages=collections.defaultdict(list)
for r in x['gsc']['current']['query_page']:pages[r['keys'][0]].append(r)
byq=collections.defaultdict(list)
for r in x['results']:byq[r['query']].append(r)
def target(q):
 if 'mayer' in q:return '/urun/limon-mayer','P1','Mevcut Mayer sayfasında Adana veri kapsamını açıklayan içerik; yerel seri tamamlanınca şehir sayfası.'
 if 'limon' in q:
  if 'adana' in q:return '/urun/limon','P0','Adana resmi veri akışını doğrula/onar; genel fiyatı Adana fiyatı diye sunma; yerel sayfa veri kapısı sonrası.'
  if 'mersin' in q or 'erdemli' in q:return '/piyasa/erdemli-limon','P0','Yerel veri akışını onar; piyasa, analiz ve ürün sayfalarının amaçlarını ve iç bağlantılarını ayır.'
  return '/urun/limon','P1','Çeşit, hal/bahçe ayrımı, kaynaklı tablo ve ilgili Mayer/yerel piyasa bağlantılarını güçlendir.'
 if 'salçalık domates' in q:return '/urun/domates-salcalik','P1','Mevcut sayfayı Rio/çeşit, hal-fabrika ayrımı ve veriye dayalı hesap örneğiyle güçlendir.'
 if 'üzüm' in q:return ('/urun/uzum' if 'kuru' not in q else ''),'P1','Taze/kuru üzüm niyetini ayır; mevcut hedefi koru, çeşide ve kaynağa bağlı açıklama ekle.'
 if 'mersin' in q:return '/hal/mersin-hal','P0','Son kaynak tarihi 22 Haziran; yeni yazıdan önce güncel resmi kaynak ve ETL sorununu çöz.'
 if 'bayrampaşa' in q or 'istanbul' in q or 'ibb' in q:return '/hal/istanbul-hal-ibb','P2','Aynı hal sayfasında Bayrampaşa fiyat listesi başlığı ve güncel satırları önce göster; kopya URL açma.'
 for c,slug in [('kocaeli','kocaeli-hal-merkez'),('konya','konya-hal'),('bursa','bursa-hal'),('ankara','ankara-hal'),('denizli','denizli-hal'),('kayseri','kayseri-hal'),('gaziantep','gaziantep-hal'),('kahramanmaraş','kahramanmaras-hal')]:
  if c in q:return '/hal/'+slug,'P2','Google mevcut hedefini kontrol et; veri tarihi, liste sırası ve sorguya uygun açıklamayı güçlendir.'
 if q=='hal fiyatları':return '/fiyatlar','P1','Ana sayfa/fiyatlar sorgu sahipliğini incele; şehir bağlantıları ve son veri tarihlerini güçlendir.'
 ps=sorted(pages.get(q,[]),key=lambda r:-r['impressions']);return (ps[0]['keys'][1].replace('https://haldefiyat.com','') if ps else ''),'P2','Mevcut açılış sayfasında niyet ve veri kapsamı incelemesi.'
qr=[]
for q,rs in byq.items():
 g=cur.get(q,{});p=prev.get(q,{});ps=sorted(pages.get(q,[]),key=lambda r:-r['impressions']);t,priority,action=target(q);ours=min([r['position'] for r in rs if r['is_ours']],default=None);best={}
 for r in rs:
  if not r['is_ours'] and (r['domain'] not in best or r['position']<best[r['domain']]['position']):best[r['domain']]=r
 ahead=[r for r in best.values() if ours is None or r['position']<ours]
 qr.append(dict(query=q,priority=priority,gsc_impressions=g.get('impressions',0),gsc_clicks=g.get('clicks',0),gsc_ctr_pct=round(g.get('ctr',0)*100,2),gsc_position=round(g.get('position',0),2),previous_position=round(p.get('position',0),2),previous_clicks=p.get('clicks',0),sample_our_position=ours,competitors_ahead_in_sample=len(ahead),google_top_page=ps[0]['keys'][1] if ps else '',target_path=t,action=action))
qr.sort(key=lambda r:(r['priority'],-r['gsc_impressions']))
domains=[];gaps=[]
for domain in sorted(set(r['domain'] for r in x['results'] if not r['is_ours'])):
 rs=[r for r in x['results'] if r['domain']==domain];best={}
 for r in rs:
  if r['query'] not in best or r['position']<best[r['query']]['position']:best[r['query']]=r
 ahead=0;imps=0
 for q,r in best.items():
  ours=next(a['sample_our_position'] for a in qr if a['query']==q);a=ours is None or r['position']<ours
  ahead+=a;imps+=cur.get(q,{}).get('impressions',0)
  gaps.append(dict(domain=domain,query=q,sample_competitor_position=r['position'],sample_our_position=ours,ahead_in_sample=a,competitor_url=r['url'],our_gsc_position=round(cur.get(q,{}).get('position',0),2),our_gsc_impressions=cur.get(q,{}).get('impressions',0),target_path=target(q)[0],action=target(q)[2]))
 kind='resmi' if any(z in domain for z in ['.bel.tr','.gov.tr','ibb.istanbul']) else 'diger'
 domains.append(dict(domain=domain,kind=kind,queries=len(best),sample_average_position=round(sum(r['position'] for r in best.values())/len(best),2),ahead_queries=ahead,our_gsc_impressions_on_shared_queries=imps,sample_url=min(rs,key=lambda r:r['position'])['url']))
domains.sort(key=lambda r:(-r['queries'],r['sample_average_position']))
extra=[]
for q,g in cur.items():
 if q not in byq and g['impressions']>=300 and g['position']>7:
  t,pr,ac=target(q);extra.append(dict(query=q,impressions=g['impressions'],clicks=g['clicks'],position=round(g['position'],2),target_path=t,action=ac))
extra.sort(key=lambda r:-r['impressions'])
wb=Workbook();wb.remove(wb.active)
for name,rows in [('Sorgular',qr),('Tum rakipler',domains),('Rakip sorgu eslesmeleri',gaps),('Izleme disi firsatlar',extra)]:
 with (P/(name.replace(' ','-')+'.csv')).open('w',newline='') as f:w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
 ws=wb.create_sheet(name);ws.append(list(rows[0]));
 for r in rows:ws.append(list(r.values()))
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='14532D')
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col[:20])+2))
notes=wb.create_sheet('Yontem');
for row in [('Alan','Tanım'),('GSC güncel','10 Ağustos–6 Eylül 2026; final data; web; tüm ülke/cihazlar'),('GSC önceki','13 Temmuz–9 Ağustos 2026; aynı 28 günlük uzunluk'),('SERP','5 Eylül 2026, run 4, Brave + 1 sayfa Yandex fallback; 30 sorgu; 534 kayıt'),('Yok','Örneklemde görünmedi; Google indeks durumu değildir.'),('Gösterim','Bizim mülkün GSC gösterimi; rakibin trafiği veya toplam arama hacmi değildir.'),('Birleştirme','Rakip-domain/sorgu başına en iyi URL; sorgular domainler arasında tekrar ettiği için domain gösterimleri toplanmaz.'),('GSC query/page','Sayfa gösterimleri query toplamına eşit olmak zorunda değildir; birden çok URL görünmesi tek başına kanibalizasyon kanıtı değildir.'),('Yorum','Rakip Google pozisyonu ölçülmedi; Google karşılaştırması diye Brave/GSC farkı hesaplanmaz.'),('Kapsam','194 domain yalnız taranan 30 sorguda; sitelerin tüm anahtar kelimelerini kapsamaz.'),('Veri sınırı','GSC anonim sorguları API tarafından dışarıda bırakılır; sıralamalar ortalamadır, kişisel anlık sonuç değildir.')]:notes.append(row)
notes.column_dimensions['A'].width=25;notes.column_dimensions['B'].width=120
wb.save(P/'rakip-seo-oncelikleri.xlsx')
(P/'computed.json').write_text(json.dumps(dict(queries=qr,domains=domains,extras=extra),ensure_ascii=False,indent=2))
print(len(qr),'queries',len(domains),'domains',len(gaps),'pairs',len(extra),'extra opportunities')
