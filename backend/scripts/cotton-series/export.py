#!/usr/bin/env python3
"""Export the validated cotton dataset, workbook and source-linked article fragment."""
import argparse, csv, html, json, math, shutil
from pathlib import Path
from collections import Counter
from datetime import date
from decimal import Decimal
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

p=argparse.ArgumentParser();p.add_argument('--data',default='data/cotton-series');args=p.parse_args();root=Path(args.data)
rows=json.loads((root/'observations.json').read_text());summary=json.loads((root/'summary.json').read_text());coverage=json.loads((root/'coverage.json').read_text());q=json.loads((root/'quarantine.json').read_text())
if summary['through']!='2026-09-07':raise ValueError('This editorial export template is fixed to 2026-09-07; update text and labels before publishing another cutoff')
public=root/'public';public.mkdir(exist_ok=True)
base='/uploads/cotton-series/'+summary['through']+'/'
for name in ['observations.csv','observations.json','coverage.json','quarantine.json','summary.json']:shutil.copy2(root/name,public/name)
fields=list(rows[0]);wb=Workbook();ws=wb.active;ws.title='Aylık kayıtlar';ws.append(fields)
number_fields={'min_price','max_price','avg_price','quantity_kg','turnover_try','derived_weighted_price'}
for r in rows:ws.append([float(r[k]) if k in number_fields and r[k] is not None else r[k] for k in fields])
ws.freeze_panes='H2';ws.auto_filter.ref=ws.dimensions
for title,content in [('Kapsam',coverage),('Karantina',[{'error':e['error'],'source_url':e.get('source_url',e.get('row',{}).get('source_url')),'raw_line':e.get('raw_line',e.get('row',{}).get('raw_line',''))} for e in q])]:
 sh=wb.create_sheet(title);cols=list(dict.fromkeys(k for r in content for k in r));sh.append(cols)
 for r in content:sh.append([r.get(k) for k in cols])
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
notes=[
 ('Kapsam',f"2013 Ocak–{summary['through']}; {len(rows)} doğrulanmış aylık sınıf/işlem/ödeme kaydı. Son ay tamamlanmamış olabilir."),
 ('Kaynaklar','Şanlıurfa Ticaret Borsası tarihli tescil bültenleri; İzmir Ticaret Borsası aylık pamuk kapanış bültenleri.'),
 ('Fiyat birimi','TRY/kg; nominal Türk lirası, enflasyondan arındırılmamış.'),
 ('avg_price','Kaynak bültenin ortalama sütunu. ŞUTB ortalama ağırlıklandırma yöntemi belirtilmediğinden varsayılmaz.'),
 ('derived_weighted_price','ŞUTB toplam tutar / toplam miktar hesabı; HaldeFiyat türetimidir, kaynak ortalaması değildir.'),
 ('Seri kimliği','Kaynak, tam ürün etiketi, kütlü/lif/linter, işlem türü, ödeme türü, fiyat türü, birim ve ortalama yöntemi aynı olmalı.'),
 ('Dönem','period_complete=false ise kapanmamış ay; tam aylarla yıllık yüzde karşılaştırmasına katılmaz.'),
 ('Eksik kayıt','Boş fiyatlar sıfır yapılmaz; kaynaksız aylar interpolasyonla doldurulmaz. Aynı etiketin randıman/mahsul yılı aynı kabul edilmez.'),
 ('Karantina','Yanlış belge, eksik fiyat, aralık dışı ortalama, geçersiz birim, tutar/miktar tutarsızlığı ve aynı boyutlu çoğul satırlar ana seriden ayrıldı.'),
 ('Yöntem değişikliği','İTB hazır peşin + ismen ortalaması ile yalnız hazır peşin ortalaması ayrı series_id ile saklanır. Eski Std.1 ve yeni renk kodları eşdeğer varsayılmaz.'),
 ('Kaynak izi','Her satırın source_url, source_sha256 ve raw_line alanları vardır. Orijinal cevaplar özel raw arşivinde korunur.'),
 ('İthal ürün','İthal preseli ve linter pamuk ayrı etiketlerdir; yerli preseli veya kütlü seriyle birleştirilmez.'),
 ('Aydın ve günlük dosyalar','Bu sürüm ŞUTB aylık ve İTB aylık kaynaklarını kapsar. Aydın OCR ve İTB günlük dosyaları bu seriye dahil değildir.'),
 ('Eski seri','Genel Pamuk günlük geçmişinden bağımsız yeni aylık seri; mevcut eksik günlük kayıtlar bu veri setinin parçası değildir.')]
sh=wb.create_sheet('Açıklama');sh.append(['Alan','Açıklama'])
for r in notes:sh.append(r)
for sh in wb:
 sh.sheet_view.showGridLines=False
 for c in sh[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='174B42')
 for col in sh.columns:
  letter=col[0].column_letter;sh.column_dimensions[letter].width=min(60,max(14,len(str(col[0].value))+3))
 sh.row_dimensions[1].height=30
wb.save(public/'pamuk-aylik-fiyat-serisi.xlsx')
# Two separate axes: kütlü and lint are not substitutes. Missing months remain gaps.
months=[];year,month=2013,1
while f'{year}-{month:02}-01'<=summary['through']:
 months.append(f'{year}-{month:02}-01');year,month=(year+1,1) if month==12 else (year,month+1)
fig,axes=plt.subplots(2,1,figsize=(12,8),sharex=True)
fig.patch.set_facecolor('white')
series_specs=[('PRESELİ PAMUK ST.1 BEYAZ','HTS','Preseli pamuk ST.1 beyaz · HTS · peşin','#155E75'),('KÜTLÜ PAMUK','HMS','Kütlü pamuk · HMS · peşin','#B45309')]
for ax,(prod,sale,label,color) in zip(axes,series_specs):
 points={r['period_start']:r for r in rows if r['source']=='sutb' and r['product']==prod and r['sale_type']==sale and r['payment_type']=='PEŞİN'}
 vals=[float(points[m]['avg_price']) if m in points and points[m]['period_complete'] else float('nan') for m in months]
 dates=[date.fromisoformat(m) for m in months]
 ax.plot(dates,vals,color=color,linewidth=2)
 for m,r in points.items():
  if not r['period_complete']:ax.scatter([date.fromisoformat(m)],[float(r['avg_price'])],facecolors='none',edgecolors=color,s=55,zorder=5)
 ax.set_title(label,loc='left',fontsize=13,pad=12);ax.set_ylabel('TL/kg');ax.set_ylim(bottom=0);ax.grid(axis='y',color='#E5E7EB',linewidth=.6);ax.spines[['top','right']].set_visible(False)
 ax.text(.01,.96,f"{len(points)} dönem · kaynak bülten ortalaması",transform=ax.transAxes,va='top',fontsize=9,color='#475569')
axes[-1].xaxis.set_major_locator(mdates.YearLocator(2));axes[-1].xaxis.set_major_formatter(mdates.DateFormatter('%Y'))
fig.suptitle('Şanlıurfa pamuk fiyatlarının aylık seyri',x=.08,ha='left',fontsize=19,y=.97)
fig.text(.08,.92,'Ocak 2013–7 Eylül 2026 · Nominal TL/kg · Kütlü ve preseli pamuk ayrı serilerdir.',fontsize=10)
fig.text(.08,.035,'Kaynak: Şanlıurfa Ticaret Borsası tescil bültenleri. Boş aylar doldurulmadı.\nİçi boş nokta: 1–7 Eylül 2026, tamamlanmamış ay. Mahsul yılı ve randıman bültende belirtilmiyor.',fontsize=9,color='#475569')
fig.subplots_adjust(top=.85,bottom=.13,left=.08,right=.98,hspace=.25);fig.savefig(public/'pamuk-aylik-seri.png',dpi=180);fig.savefig(public/'pamuk-aylik-seri.svg');plt.close(fig)
# Same exact series, same completed calendar month for year-on-year comparisons.
lookup={(r['series_id'],r['period_start']):r for r in rows if r['period_complete']}
comparisons=[]
for r in rows:
 if not r['period_complete']:continue
 prev=lookup.get((r['series_id'],str(int(r['period_start'][:4])-1)+r['period_start'][4:]))
 if prev:comparisons.append(dict(series_id=r['series_id'],source=r['source'],product=r['product'],sale_type=r['sale_type'],payment_type=r['payment_type'],period_start=r['period_start'],previous_period_start=prev['period_start'],avg_price=r['avg_price'],previous_avg_price=prev['avg_price'],change_pct=str(((Decimal(r['avg_price'])/Decimal(prev['avg_price'])-1)*100).quantize(Decimal('.01'))),source_url=r['source_url'],previous_source_url=prev['source_url']))
(root/'comparisons.json').write_text(json.dumps(comparisons,ensure_ascii=False,indent=2));shutil.copy2(root/'comparisons.json',public/'comparisons.json')
with (public/'comparisons.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(comparisons[0]));w.writeheader();w.writerows(comparisons)
md='# Pamuk aylık fiyat serisi\n\n'+''.join(f'## {k}\n\n{v}\n\n' for k,v in notes)
(root/'README.md').write_text(md);(public/'README.txt').write_text(md)
# Source-backed data sheet; no separate analytical aggregation hidden behind the table.
fmt=lambda v:'—' if v is None else f'{float(v):,.2f}'.replace(',','X').replace('.',',').replace('X','.')
primary=[r for r in rows if r['source']=='sutb' and r['product']=='PRESELİ PAMUK ST.1 BEYAZ' and r['sale_type']=='HTS' and r['payment_type']=='PEŞİN']
yoy=next(x for x in comparisons if x['series_id']==primary[0]['series_id'] and x['period_start']=='2026-08-01')
body=f'''<p><strong>Pamukta karşılaştırmayı resmî borsa kayıtlarıyla yapıyoruz.</strong> Ocak 2013’ten başlayan arşiv taramasında {summary['bulletins']} bülten sorgusu ve dosyası işlendi; {format(len(rows), ',').replace(',', '.')} geçerli dönem–ürün–işlem–ödeme kaydı ayrı fiyat serileri halinde derlendi. Kaynaklar Şanlıurfa Ticaret Borsası’nın tescil bültenleri ve İzmir Ticaret Borsası’nın aylık pamuk bültenleri.</p>
<h2>Ağustos 2026: aynı başlıkta yıllık değişim %{fmt(yoy['change_pct'])}</h2>
<p>Şanlıurfa’da <strong>preseli pamuk ST.1 beyaz, HTS, peşin</strong> satırının bülten ortalaması Ağustos 2025’te <strong>{fmt(yoy['previous_avg_price'])} TL/kg</strong>, Ağustos 2026’da <strong>{fmt(yoy['avg_price'])} TL/kg</strong>. İki tamamlanmış ay arasındaki nominal değişim <strong>%{fmt(yoy['change_pct'])}</strong>. Hesap: (2026 değeri / 2025 değeri − 1) × 100. Bu karşılaştırma aynı borsa, ürün etiketi ve ödeme türüne dayanır; mahsul yılı, randıman ve parti kalitesinin aynı olduğunu göstermez. <a href="{html.escape(yoy['previous_source_url'],quote=True)}">Ağustos 2025 bülteni</a> · <a href="{html.escape(yoy['source_url'],quote=True)}">Ağustos 2026 bülteni</a>.</p>
<h2>2013’ten bugüne aylık fiyatlar</h2>
<p>Grafikte kütlü ve preseli pamuk ayrı gösteriliyor. Boş aylar sıfır fiyat değildir; çizgiyle doldurulmadı. Eylül 2026 henüz tamamlanmadığı için bu ayı tam Eylül 2025 ile kıyaslamıyoruz. Fiyatlar nominal TL/kg’dır; enflasyondan arındırılmamıştır.</p>
<p><img src="{base}pamuk-aylik-seri.png" alt="Şanlıurfa Ticaret Borsası kütlü ve preseli pamuk ayrı aylık fiyat serileri, 2013–2026" /></p>
<h2>Veriyi indir ve kaynağı kontrol et</h2>
<p><a href="{base}pamuk-aylik-fiyat-serisi.xlsx">Excel: fiyatlar, kapsam, karantina ve yöntem</a> · <a href="{base}observations.csv">Tüm fiyat kayıtları CSV</a> · <a href="{base}comparisons.csv">Tam ay, aynı seri yıllık karşılaştırmaları CSV</a> · <a href="{base}coverage.json">Bülten kapsamı</a> · <a href="{base}quarantine.json">Seriye alınmayan kayıtlar</a>.</p>
<h2>Veri kapsamı ve sınırlar</h2>
<p>Şanlıurfa’da aynı ST.1 beyaz / HTS / peşin başlığı {len(primary)} aylık dönemde izlendi. Kütlü pamuk, ithal preseli pamuk, linter, farklı kalite ve vadeli işlemler birbirine eklenmedi. İzmir’de eski standart adları ile yeni renk kodları otomatik eşlenmedi; ortalama hesap yönteminin değiştiği seriler ayrıldı.</p>
<p>{len(q)} satır/belge kaydı; eksik fiyat, tutarsız değer, birim sorunu, aynı boyutlarda çoğul satır veya yanlış belge nedeniyle ana seriye alınmadı. Örneğin Kasım 2017 İzmir pamuk bağlantısı bir dünya emtia raporuna açılıyor. Arşivde bağlantı bulunması, geçerli pamuk fiyatının bulunduğu anlamına gelmiyor.</p>
<p>Borsa ortalaması ile toplam işlem tutarı / miktardan türetilen fiyat ayrı sütunlarda saklanır. İşlem sayısı ve miktarı düşük dönemlerin fiyatları geniş piyasa geneli diye yorumlanmamalıdır. Fiyatlar Türkiye geneli üretici ortalaması veya doğrudan kârlılık ölçüsü değildir.</p>
<h2>Şanlıurfa ST.1 beyaz, peşin: aylık kayıt tablosu</h2>
<p>Aşağıdaki tablo aynı ürün ve işlem başlığını izler. Her satırın kaynağı açılabilir; ortalama, bültende yayımlanan değerdir. Diğer ürün ve kalite serileri Excel/CSV dosyasında ayrı kimliklerle bulunur.</p>
<table><thead><tr><th>Dönem</th><th>En düşük TL/kg</th><th>Ortalama TL/kg</th><th>En yüksek TL/kg</th><th>İşlem</th><th>Miktar kg</th><th>Kaynak</th></tr></thead><tbody>'''
for r in sorted(primary,key=lambda x:x['period_start'],reverse=True):
 label=r['period_start'][:7]+(' (1–7, kısmi)' if not r['period_complete'] else '')
 body+=f'<tr><td>{label}</td><td>{fmt(r["min_price"])}</td><td>{fmt(r["avg_price"])}</td><td>{fmt(r["max_price"])}</td><td>{r["transaction_count"]}</td><td>{fmt(r["quantity_kg"])}</td><td><a href="{html.escape(r["source_url"],quote=True)}">Bülten</a></td></tr>'
body+='</tbody></table><p>Veri kesim tarihi: 7 Eylül 2026. Kaynak arşivleri: <a href="https://itb.org.tr/AylikBultenler/5-pamuk-bulteni">İzmir Ticaret Borsası</a> ve <a href="https://uye.sutb.org.tr:3333/bultenweb">Şanlıurfa Ticaret Borsası</a>. Bu sürüm aylık seridir; günlük İzmir dosyaları ve Aydın arşivi kapsam dışındadır.</p>'
(root/'article.html').write_text(body)
(root/'publication.json').write_text(json.dumps(dict(slug='pamuk-fiyatlari-gecmisi-2013-2026-borsa-serisi',title='Pamuk Fiyatları 2013–2026: Resmî Borsa Verileriyle Aylık Seri',summary=f"Şanlıurfa ve İzmir ticaret borsalarından {len(rows)} doğrulanmış kayıt. Kütlü ve lif pamuk ayrı serilerde; kaynak bağlantıları ve indirilebilir Excel/CSV.",content=body,metaTitle='Pamuk Fiyatları 2013–2026 | Resmî Aylık Borsa Serisi',metaDescription='2013–2026 pamuk fiyat geçmişi: Şanlıurfa ve İzmir resmî borsa kayıtları, kalite ve ödeme türüne göre aylık seriler. Kaynaklı Excel ve CSV indir.',ogImage=base+'pamuk-aylik-seri.png',tags=['pamuk','fiyat geçmişi','borsa verileri','aylık seri','Şanlıurfa','İzmir']),ensure_ascii=False,indent=2))
print(json.dumps({'rows':len(rows),'comparisons':len(comparisons),'primary_periods':len(primary),'yoy_august_2026':yoy},ensure_ascii=False))
